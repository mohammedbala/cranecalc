"""Extract W-shape US nominal dimensions/properties from the pinned AISC v16 workbook.
Usage: python scripts/extract-aisc.py path/to/aisc-shapes-database-v16.0.xlsx
Requires openpyxl. Output uses original US values, with no intermediate SI rounding.
"""
import hashlib, json, pathlib, sys
import openpyxl
path = pathlib.Path(sys.argv[1])
expected = 'dec6202f8c2055bdefee9e13cb3d38e9110664c61dcf91d09604296ef2dcf660'
assert hashlib.sha256(path.read_bytes()).hexdigest() == expected, 'Unexpected source workbook'
book = openpyxl.load_workbook(path, read_only=True, data_only=True)
sheet = book['Database v16.0']
headers = next(sheet.iter_rows(values_only=True))
fields = ['W', 'd', 'bf', 'tw', 'tf', 'A', 'Ix', 'Iy', 'Sx', 'Sy', 'Zx', 'Zy', 'J', 'Cw', 'kdes']
indices = {key: headers.index(key) for key in fields}
shapes = []
for row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    if values[0] != 'W': continue
    shape = {'name': values[1], 'row': row}
    for key, index in indices.items():
        value = values[index]
        assert isinstance(value, (int, float)) and value > 0, (shape['name'], key, value)
        shape['weight' if key == 'W' else key] = value
    shapes.append(shape)
assert len({s['name'] for s in shapes}) == len(shapes)
out = pathlib.Path(__file__).resolve().parents[1] / 'src/data/aiscWShapes.json'
out.write_text(json.dumps(shapes, separators=(',', ':')) + '\n')
print(f'Extracted {len(shapes)} W-shapes to {out}')
