"""Extract unrounded US C/MC catalogue values from the pinned AISC workbook."""
import hashlib,json,pathlib,sys
import openpyxl
path=pathlib.Path(sys.argv[1])
assert hashlib.sha256(path.read_bytes()).hexdigest()=='dec6202f8c2055bdefee9e13cb3d38e9110664c61dcf91d09604296ef2dcf660'
sheet=openpyxl.load_workbook(path,read_only=True,data_only=True)['Database v16.0']
headers=next(sheet.iter_rows(values_only=True))
fields=['W','A','d','bf','tw','tf','x','Ix','Iy','Sx','Sy','Zx','Zy','J','Cw','kdes']
out=[]
for n,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
 if values[0] not in ['C','MC']:continue
 shape={'name':values[1],'row':n}
 for key in fields:
  v=values[headers.index(key)];assert isinstance(v,(int,float)) and v>0,(shape['name'],key,v)
  shape['weight' if key=='W' else key]=v
 out.append(shape)
assert len({s['name'] for s in out})==len(out)
target=pathlib.Path(__file__).resolve().parents[1]/'src/data/aiscChannels.json'
target.write_text(json.dumps(out,separators=(',',':'))+'\n')
print(f'Extracted {len(out)} channels')
